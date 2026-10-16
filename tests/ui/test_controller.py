"""Tests for the GUI controller wiring (offscreen Qt)."""
import os
import threading
import time
import types

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest  # noqa: E402
from PySide6.QtCore import QCoreApplication  # noqa: E402
from PySide6.QtWidgets import QApplication, QMessageBox  # noqa: E402

from store_sync.core.config import (  # noqa: E402
    BindingConfig,
    ColumnBinding,
    JobConfig,
    Schedule,
    SourceConfig,
    StoreConfig,
    WriteConfig,
)
from store_sync.core.registry import JobEntry  # noqa: E402
from store_sync.gui import Controller  # noqa: E402
from store_sync.ui.main_window import MainWindow  # noqa: E402

_app = QApplication.instance() or QApplication([])


def a_job(name="a"):
    return JobConfig(
        name,
        StoreConfig(17620, "No"),
        SourceConfig("SURV"),
        BindingConfig(
            ColumnBinding("SbjNum", "SbjNum"),
            [ColumnBinding("T_QPM1_1", "I7_1_2")],
        ),
        WriteConfig(dry_run=True),
    )


def test_save_lists_on_dashboard(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    assert c.list_entries()[0].name == "a"
    w.show_dashboard()
    assert w.dashboard.model.rowCount() == 1
    c.stop()
    w.close()


def test_run_now_enqueues(tmp_path):
    w = MainWindow()
    calls = []
    c = Controller(
        w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: calls.append(name)
    )
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    c.start()
    c.run_now("a")
    time.sleep(0.2)
    assert calls == ["a"]
    c.stop()
    w.close()


def test_save_twice_updates_not_duplicates(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    assert w.dashboard.model.rowCount() == 1
    c.save("a", a_job(), schedule=Schedule("hour", 1), enabled=True)
    assert len(c.list_entries()) == 1
    assert w.dashboard.model.rowCount() == 1
    entry = c.list_entries()[0]
    assert entry.schedule == Schedule("hour", 1)
    assert entry.enabled is True
    c.stop()
    w.close()


def test_delete_removes_and_refreshes(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    assert w.dashboard.model.rowCount() == 1
    c.delete("a")
    assert c.list_entries() == []
    assert w.dashboard.model.rowCount() == 0
    c.stop()
    w.close()


def test_stop_is_idempotent(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.start()
    c.stop()
    c.stop()  # must not raise the second time
    w.close()


def test_save_defaults_blank_match_store_col(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    base = a_job()
    blank = JobConfig(
        base.job,
        base.store,
        base.source,
        BindingConfig(ColumnBinding("SbjNum", ""), base.binding.columns),
        base.write,
    )
    c.save("a", blank, schedule=Schedule(""), enabled=False)
    assert c.list_entries()[0].match_store_col == "SbjNum"
    c.stop()
    w.close()


def test_worker_thread_db_access(tmp_path):
    """The queue worker must touch Registry/State without a thread error."""
    w = MainWindow()
    results = []

    def run_fn(name, stop):
        results.append(c.registry.get(name))
        results.append(c.state.start_run(job=name, dry_run=True))

    c = Controller(w, db_path=tmp_path / "s.db", run_fn=run_fn)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    c.start()
    c.run_now("a")
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and not results:
        time.sleep(0.02)
    c.stop()
    w.close()

    assert len(results) == 2
    assert results[0].job == "a"
    assert isinstance(results[1], int)
    run = c_state_get(tmp_path / "s.db", results[1])
    assert run is not None and run["job"] == "a"


def test_registry_from_plain_thread(tmp_path):
    """`list_entries` from a non-GUI thread must not raise a thread error."""
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)

    out = {}

    def worker():
        try:
            out["entries"] = [e.name for e in c.registry.list_entries()]
        except Exception as exc:  # noqa: BLE001
            out["error"] = exc

    th = threading.Thread(target=worker)
    th.start()
    th.join(3)
    c.stop()
    w.close()

    assert "error" not in out, out["error"]
    assert out["entries"] == ["a"]


def test_save_bad_config_reports_error(tmp_path):
    """A non-numeric store id must not raise out of the save slot."""
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._detail_name = "a"
    w.show_job("a")
    w.job_detail.store_id_edit.setText("not-a-number")
    c._save_from_detail()  # must not raise
    assert c.list_entries() == []
    assert "error" in w.status_label.text().lower()
    assert "#f87171" in w.status_label.styleSheet()
    c.stop()
    w.close()


def test_client_missing_env_names_the_vars(tmp_path, monkeypatch):
    """`_client()` must say WHICH vars are missing and point at `.env`."""
    for var in ("STG_SERVER", "STG_USER", "STG_PASSWORD"):
        monkeypatch.delenv(var, raising=False)
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    with pytest.raises(ValueError, match=r"STG_SERVER.*STG_USER.*STG_PASSWORD.*\.env"):
        c._client()
    c.stop()
    w.close()


def test_client_partial_env_names_only_missing(tmp_path, monkeypatch):
    """Only the missing vars are named; present ones are not required again."""
    monkeypatch.setenv("STG_SERVER", "https://stg.example")
    monkeypatch.delenv("STG_USER", raising=False)
    monkeypatch.delenv("STG_PASSWORD", raising=False)
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    with pytest.raises(ValueError, match="STG_USER/STG_PASSWORD"):
        c._client()
    msg = None
    try:
        c._client()
    except ValueError as exc:
        msg = str(exc)
    assert "STG_SERVER" not in msg  # set → not listed
    c.stop()
    w.close()


def test_export_missing_env_surfaces_actionable_error(tmp_path, monkeypatch):
    """Export with no .env must toast the missing-vars message, not a traceback."""
    for var in ("STG_SERVER", "STG_USER", "STG_PASSWORD"):
        monkeypatch.delenv(var, raising=False)
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._detail_name = "a"
    w.job_detail.load(_entry(), a_job())
    w.show()  # toasts are child overlays → visible only with the window up
    monkeypatch.setattr(
        "PySide6.QtWidgets.QFileDialog.getSaveFileName",
        staticmethod(lambda *a, **k: ("m.xlsx", "")),
    )

    c._export_mapping()
    _drain(c, "_mapping_thread")
    from store_sync.ui.toast import Toast

    toasts = [t for t in w.findChildren(Toast) if t.isVisible()]
    assert toasts, "expected a failure toast"
    assert "STG_SERVER" in toasts[0].text()
    assert ".env" in toasts[0].text()
    assert "error" in w.status_label.text().lower()
    c.stop()
    w.close()


def test_connections_refused_after_stop(tmp_path):
    """A lazy connection request after `stop()` must not leak a tracked conn."""
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.start()
    c.stop()
    for factory in (c._state, c._registry):
        try:
            factory()
        except RuntimeError:
            pass
        else:  # pragma: no cover - a returned conn would be a leak
            raise AssertionError("factory returned a connection after stop()")
    assert c._conns == []
    w.close()


def c_state_get(path, run_id):
    """Read a run row with a fresh connection (no shared-thread issues)."""
    from store_sync.core.state import State

    st = State(path)
    try:
        return st.get_run(run_id)
    finally:
        st.close()


# -- Important 1: history / conflicts / run detail --------------------------

def _seed_run(c, *, job="a", status="ok", counts=None, conflict=None, dry_run=True):
    """Record one finished run (+ optional conflict) via the controller's State."""
    run_id = c.state.start_run(job=job, dry_run=dry_run)
    if conflict is not None:
        c.state.record_conflict(run_id, match_value=conflict[0], reason=conflict[1])
    c.state.finish_run(run_id, status=status, counts=counts or {"add": 1})
    return run_id


def test_history_marks_dry_run_runs(tmp_path):
    """A dry run's counts must read '(dry run)' — never look like real writes."""
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    _seed_run(c, dry_run=True, counts={"add": 296, "delete": 1})
    _seed_run(c, dry_run=False, counts={"add": 2})
    c._open_job("a")
    rows = w.job_detail.history_model._rows
    assert "(dry run)" in rows[1]["counts"]   # newest first: live run is [0]
    assert "(dry run)" not in rows[0]["counts"]
    c.stop()
    w.close()


def test_finished_run_toast_flags_dry_run(tmp_path):
    """The 'finished' toast must say dry run so counts aren't mistaken for writes."""
    from store_sync.ui.toast import Toast

    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    _seed_run(c, dry_run=True, counts={"add": 296, "delete": 1})
    w.show()  # child-overlay toasts need a visible window
    c._active_prev = "a"  # pretend a run was in flight
    from store_sync.core.run_queue import QueueState

    c._on_queue_changed(QueueState(active=None, pending=[]))
    toasts = [t for t in w.findChildren(Toast) if t.isVisible()]
    assert toasts, "expected a finished toast"
    assert "dry run" in toasts[0].text().lower()
    c.stop()
    w.close()


def test_run_progress_reaches_job_detail(tmp_path):
    """`run_sync` progress events must drive the JobDetail timeline panel."""
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    c._open_job("a")
    assert w.job_detail.progress_panel.isHidden()
    c._on_run_progress({"stage": "start", "job": "a", "dry_run": True})
    c._on_run_progress({"stage": "ids", "total": 296})
    QCoreApplication.processEvents()
    assert not w.job_detail.progress_panel.isHidden()
    assert "296" in w.job_detail.progress_text.toPlainText()
    c.stop()
    w.close()


def test_run_progress_updates_header_status(tmp_path):
    """The header status line must show the current stage (visible off-page)."""
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    c._on_run_progress({"stage": "pull", "total": 296})
    assert "296" in w.status_label.text()
    assert "#fbbf24" in w.status_label.styleSheet()  # warn while running
    c._on_run_progress({"stage": "done", "status": "ok", "counts": {}})
    assert "ok" in w.status_label.text().lower()
    c.stop()
    w.close()


def test_progress_does_not_bleed_across_jobs(tmp_path):
    """Job A's run must not render into job B's open detail panel."""
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    c.save("b", a_job("b"), schedule=Schedule(""), enabled=False)
    c._open_job("b")
    assert w.job_detail.progress_panel.isHidden()
    c._on_run_progress({"stage": "start", "job": "a", "dry_run": False})
    c._on_run_progress({"stage": "ids", "job": "a", "total": 296})
    QCoreApplication.processEvents()
    # B is open and idle → A's events must not fill B's panel.
    assert w.job_detail.progress_panel.isHidden()
    assert w.job_detail.progress_text.toPlainText() == ""
    c.stop()
    w.close()


def test_open_running_job_replays_its_buffered_log(tmp_path):
    """Opening a job that is mid-run shows the events already emitted for it."""
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    c._on_run_progress({"stage": "start", "job": "a", "dry_run": False})
    c._on_run_progress({"stage": "ids", "job": "a", "total": 296})
    c._on_run_progress({"stage": "pull", "job": "a", "total": 296})
    QCoreApplication.processEvents()
    c._open_job("a")
    QCoreApplication.processEvents()
    text = w.job_detail.progress_text.toPlainText()
    assert "296" in text
    assert not w.job_detail.progress_panel.isHidden()
    c.stop()
    w.close()


def test_open_idle_job_clears_stale_progress(tmp_path):
    """Opening a job with no buffered events shows an empty (hidden) panel."""
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    c.save("b", a_job("b"), schedule=Schedule(""), enabled=False)
    c._open_job("a")
    c._on_run_progress({"stage": "start", "job": "a", "dry_run": False})
    c._on_run_progress({"stage": "ids", "job": "a", "total": 296})
    QCoreApplication.processEvents()
    c._open_job("b")
    QCoreApplication.processEvents()
    assert w.job_detail.progress_text.toPlainText() == ""
    assert w.job_detail.progress_panel.isHidden()
    c.stop()
    w.close()


def test_open_job_fills_history_and_conflicts(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    _seed_run(c, conflict=("7", "duplicate key"))
    c._open_job("a")
    assert w.job_detail.history_model.rowCount() >= 1
    assert w.job_detail.conflicts_model.rowCount() == 1
    row = w.job_detail.history_model._rows[0]
    assert {"time", "status", "counts"} <= set(row)
    c.stop()
    w.close()


def test_open_job_without_runs_is_empty(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    c._open_job("a")
    assert w.job_detail.history_model.rowCount() == 0
    assert w.job_detail.conflicts_model.rowCount() == 0
    c.stop()
    w.close()


def test_activate_history_row_loads_run_detail(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    run_id = _seed_run(c, status="ok", counts={"add": 2},
                       conflict=("9", "dup"))
    c._open_job("a")
    seen = []
    w.run_activated.connect(seen.append)
    w.job_detail.history_view.doubleClicked.emit(
        w.job_detail.history_model.index(0, 0)
    )
    assert seen == [run_id]
    assert w.stack.currentWidget() is w.run_detail
    assert f"Run #{run_id}" in w.run_detail.summary_label.text()
    assert w.run_detail.conflicts_model.rowCount() == 1
    c.stop()
    w.close()


def test_run_activated_noop_for_missing_run(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._show_run(999999)  # unknown id must not raise or switch pages
    assert w.stack.currentWidget() is w.dashboard
    c.stop()
    w.close()


# -- Important 2: schedules overview ----------------------------------------

def test_main_window_relays_schedules_requested():
    w = MainWindow()
    seen = []
    w.schedules_requested.connect(lambda: seen.append(1))
    w.dashboard.schedules_requested.emit()
    assert seen == [1]
    w.close()


def test_controller_show_schedules_builds_rows(tmp_path, monkeypatch):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule("day", 1, "03:00"), enabled=True)
    c.save("b", a_job("b"), schedule=Schedule(""), enabled=False)
    shown = {}

    class FakeOverview:
        def __init__(self, rows, parent=None):
            shown["rows"] = rows

        def exec(self):
            return 0

    monkeypatch.setattr("store_sync.gui.ScheduleOverview", FakeOverview, raising=False)
    c._show_schedules()  # must not raise
    rows = shown["rows"]
    assert len(rows) == 2
    assert {"name", "schedule", "enabled"} <= set(rows[0])
    c.stop()
    w.close()


# -- Important 3: delete confirmation ---------------------------------------

def test_confirm_delete_cancel_keeps_job(tmp_path, monkeypatch):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    monkeypatch.setattr(
        "store_sync.gui.QMessageBox.question",
        staticmethod(lambda *a, **k: QMessageBox.No),
    )
    c._confirm_delete("a")
    assert [e.name for e in c.list_entries()] == ["a"]
    c.stop()
    w.close()


def test_confirm_delete_yes_removes_job(tmp_path, monkeypatch):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    monkeypatch.setattr(
        "store_sync.gui.QMessageBox.question",
        staticmethod(lambda *a, **k: QMessageBox.Yes),
    )
    c._confirm_delete("a")
    assert c.list_entries() == []
    c.stop()
    w.close()


# -- Important 4: off-thread preflight / mapping ----------------------------

class _BlockingClient:
    """Fake STG client whose fetches block until `release` is set."""

    def __init__(self, release, calls):
        self.session = types.SimpleNamespace(close=lambda: calls.append("closed"))
        self.release = release
        self.calls = calls

    def get_store(self, store_id):
        self.release.wait(3)
        self.calls.append(("get_store", store_id))
        return [{"No": 1}]

    def get_survey_xml(self, survey_id):
        self.release.wait(3)
        self.calls.append(("get_survey_xml", survey_id))
        return "<Survey/>"


def _entry():
    return JobEntry(
        name="a", store_id=17620, id_column="No", survey_id="SURV",
        match_source_var="SbjNum", match_store_col="SbjNum",
        dry_run=True,
        enabled=False, schedule=Schedule(""), column_count=1, last_status=None,
        last_finished_at=None,
    )


def test_mapping_import_off_thread(tmp_path, monkeypatch):
    import store_sync.gui as gui

    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._detail_name = "a"
    w.job_detail.load(_entry(), a_job())

    release = threading.Event()
    calls: list = []
    client = _BlockingClient(release, calls)
    monkeypatch.setattr(c, "_client", lambda: client)
    monkeypatch.setattr(
        "store_sync.gui.parse_mapping",
        lambda path, store_cols, source_vars: (
            [ColumnBinding("SbjNum", "No")], []
        ),
    )
    monkeypatch.setattr(
        "PySide6.QtWidgets.QFileDialog.getOpenFileName",
        staticmethod(lambda *a, **k: ("m.xlsx", "")),
    )

    c._import_mapping()
    assert calls == []  # pool fetch has not happened on the GUI thread
    release.set()
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and c._mapping_thread.isRunning():
        QCoreApplication.processEvents()
        time.sleep(0.01)
    QCoreApplication.processEvents()
    assert w.job_detail.binding_model.rowCount() == 1
    assert "closed" in calls
    c.stop()
    assert c._mapping_thread is None
    w.close()


def test_mapping_export_off_thread(tmp_path, monkeypatch):
    import store_sync.gui as gui

    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._detail_name = "a"
    w.job_detail.load(_entry(), a_job())
    release = threading.Event()
    calls: list = []
    client = _BlockingClient(release, calls)
    monkeypatch.setattr(c, "_client", lambda: client)
    written: list = []
    monkeypatch.setattr(
        "store_sync.gui.export_template",
        lambda path, store_cols: written.append((path, list(store_cols))),
    )
    monkeypatch.setattr(
        "PySide6.QtWidgets.QFileDialog.getSaveFileName",
        staticmethod(lambda *a, **k: ("m.xlsx", "")),
    )

    c._export_mapping()
    assert calls == []
    release.set()
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and c._mapping_thread.isRunning():
        QCoreApplication.processEvents()
        time.sleep(0.01)
    QCoreApplication.processEvents()
    assert written and written[0][0] == "m.xlsx"
    c.stop()
    w.close()


def test_import_and_export_never_pull_survey_xml(tmp_path, monkeypatch):
    """The fast mapping paths must not download the (huge) survey XML.

    The var list only exists in the survey XML (30 s+ for large surveys);
    import defers source-var validation to preflight and export never needed
    it, so neither may call `get_survey_xml`.
    """
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._detail_name = "a"
    w.job_detail.load(_entry(), a_job())
    rows = [{"No": 1, "SbjNum": "9"}]
    calls: list = []

    class Client(_OrderedClient):
        def get_store(self, store_id):
            calls.append("get_store")
            return rows

        def get_survey_xml(self, survey_id):
            calls.append("get_survey_xml")  # must never happen
            return "<Survey/>"

    monkeypatch.setattr(c, "_client", lambda: Client(rows))
    xlsx = str(tmp_path / "m.xlsx")
    monkeypatch.setattr(
        "PySide6.QtWidgets.QFileDialog.getSaveFileName",
        staticmethod(lambda *a, **k: (xlsx, "")),
    )
    monkeypatch.setattr(
        "PySide6.QtWidgets.QFileDialog.getOpenFileName",
        staticmethod(lambda *a, **k: (xlsx, "")),
    )

    # Export (writes the file), fill its source vars in, then import.
    c._export_mapping()
    _drain(c, "_mapping_thread")
    assert (tmp_path / "m.xlsx").exists()
    from openpyxl import load_workbook

    wb = load_workbook(xlsx)
    ws = wb.active
    ws["B2"] = "T_QPM1_1"  # No ← T_QPM1_1
    ws["B3"] = "SbjNum"    # SbjNum ← SbjNum
    wb.save(xlsx)

    c._import_mapping()
    _drain(c, "_mapping_thread")

    assert "get_survey_xml" not in calls
    assert calls.count("get_store") == 2  # once per path
    # Import loaded the filled template into the binding grid.
    assert w.job_detail.binding_model.rowCount() == 2
    c.stop()
    w.close()


def test_export_error_surfaces_reason_not_generic(tmp_path, monkeypatch):
    """A worker failure must show the cause (e.g. "server is required").

    Regression: with no .env the export worker died on
    ValueError("server is required") and the operator saw only a generic
    "export error" (or a stuck "exporting...") with no actionable reason.
    """
    for var in ("STG_SERVER", "STG_USER", "STG_PASSWORD"):
        monkeypatch.delenv(var, raising=False)
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._detail_name = "a"
    w.job_detail.load(_entry(), a_job())
    monkeypatch.setattr(
        "PySide6.QtWidgets.QFileDialog.getSaveFileName",
        staticmethod(lambda *a, **k: ("m.xlsx", "")),
    )

    c._export_mapping()
    _drain(c, "_mapping_thread")

    status = w.status_label.text().lower()
    assert "export error" in status
    assert "stg_server" in status  # the actionable reason, not a bare label
    assert "exporting" not in status  # not stuck on the transient state
    c.stop()
    w.close()


class _OrderedClient:
    """Fake client whose store rows use non-alphabetical key order."""

    def __init__(self, rows):
        self.session = types.SimpleNamespace(close=lambda: None)
        self._rows = rows

    def get_store(self, store_id):
        return self._rows

    def get_survey_xml(self, survey_id):
        return "<Survey/>"


def test_store_pool_keeps_store_column_order(tmp_path, monkeypatch):
    """Column order must follow the store's own row order, not alphabetical."""
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._detail_name = "a"
    w.job_detail.load(_entry(), a_job())
    rows = [
        {"Zebra": 1, "Alpha": 2, "mid": 3},
        {"Beta": 4},  # extra column the first row lacked → appended last
    ]
    monkeypatch.setattr(c, "_client", lambda: _OrderedClient(rows))

    store_cols = c._store_pool(a_job())
    assert store_cols == ["Zebra", "Alpha", "mid", "Beta"]
    c.stop()
    w.close()


def test_mapping_export_empty_store_reports_error(tmp_path, monkeypatch):
    """An empty store must fail the export loudly, not write a header-only file."""
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._detail_name = "a"
    w.job_detail.load(_entry(), a_job())
    monkeypatch.setattr(c, "_client", lambda: _OrderedClient([]))
    written: list = []
    monkeypatch.setattr(
        "store_sync.gui.export_template",
        lambda path, store_cols: written.append((path, list(store_cols))),
    )
    monkeypatch.setattr(
        "PySide6.QtWidgets.QFileDialog.getSaveFileName",
        staticmethod(lambda *a, **k: ("m.xlsx", "")),
    )

    c._export_mapping()
    deadline = time.monotonic() + 3
    while time.monotonic() < deadline and (
        c._mapping_thread is not None and c._mapping_thread.isRunning()
    ):
        QCoreApplication.processEvents()
        time.sleep(0.01)
    QCoreApplication.processEvents()
    assert written == []  # no template written for an empty store
    assert "error" in w.status_label.text().lower()
    assert "#f87171" in w.status_label.styleSheet()
    c.stop()
    w.close()


# -- Important 4 (round 2): worker ownership / re-entrancy / GUI delivery ----

def _arm_pairing(c, w, monkeypatch, release, calls):
    """Wire a blocking pairing worker onto `c`."""
    c._detail_name = "a"
    w.job_detail.load(_entry(), a_job())
    monkeypatch.setattr(c, "_client", lambda: _BlockingClient(release, calls))


def _drain(c, attr, timeout=3.0):
    """Spin the event loop until the `attr` slot worker stops running."""
    deadline = time.monotonic() + timeout
    worker = getattr(c, attr)
    while (
        time.monotonic() < deadline
        and worker is not None
        and worker.isRunning()
    ):
        QCoreApplication.processEvents()
        time.sleep(0.01)
    QCoreApplication.processEvents()


def test_stop_midflight_joins_unbounded_without_abort(tmp_path, monkeypatch):
    """`stop()` must join a running worker (never drop its last ref → abort)."""
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    release = threading.Event()
    calls: list = []
    _arm_pairing(c, w, monkeypatch, release, calls)

    c._pair_survey()
    worker = c._pair_thread
    assert worker is not None and worker.isRunning()
    # The strong-ref container keeps the running QThread alive (invariant).
    assert worker in c._workers

    done = {}

    def do_stop():
        c.stop()
        done["stopped"] = True

    th = threading.Thread(target=do_stop)
    th.start()
    time.sleep(0.2)
    # Unbounded join: still blocked on the in-flight request, not abandoned.
    assert "stopped" not in done
    assert c._pair_thread is worker

    release.set()  # let the client's own call return
    th.join(5)
    assert done.get("stopped") is True
    assert not worker.isRunning()
    assert c._pair_thread is None
    assert c._workers == set()
    w.close()


def test_reentrant_pairing_does_not_block_gui(tmp_path, monkeypatch):
    """A second click while in flight returns at once with a busy status."""
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    release = threading.Event()
    calls: list = []
    _arm_pairing(c, w, monkeypatch, release, calls)

    c._pair_survey()
    first = c._pair_thread
    assert first is not None and first.isRunning()

    start = time.monotonic()
    c._pair_survey()  # re-entrant click on the GUI thread
    elapsed = time.monotonic() - start
    assert elapsed < 0.5  # did not join/block on the in-flight fetch
    assert "already running" in w.status_label.text().lower()
    assert c._pair_thread is first  # not joined + replaced
    assert first.isRunning()

    release.set()
    _drain(c, "_pair_thread")
    c.stop()
    w.close()


def test_pairing_result_delivered_on_gui_thread(tmp_path, monkeypatch):
    """The worker's `ok` slot must run on the GUI (main) thread."""
    import store_sync.core.fetcher as fetcher

    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._detail_name = "a"
    w.job_detail.load(_entry(), a_job())

    ready = threading.Event()
    ready.set()  # non-blocking fake client
    monkeypatch.setattr(c, "_client", lambda: _BlockingClient(ready, []))
    monkeypatch.setattr(
        fetcher, "parse_survey_index", lambda xml: {"T_QPM1_1": "QPM1"}
    )
    monkeypatch.setattr(
        fetcher, "pair_columns",
        lambda vars_, index: ({v: "QPM1" for v in vars_}, []),
    )

    seen: dict = {}
    main_ident = threading.get_ident()

    def record(result):
        seen["ident"] = threading.get_ident()
        Controller._on_pair_ok(c, result)

    monkeypatch.setattr(c, "_on_pair_ok", record)

    c._pair_survey()
    _drain(c, "_pair_thread")

    assert seen.get("ident") == main_ident  # delivered on the GUI thread
    assert "paired" in w.status_label.text().lower()
    c.stop()
    w.close()


# -- update proposals: apply / dismiss (spec §6) ----------------------------

class _ProposalClient:
    """Fake client whose store rows back the CAS update apply."""

    def __init__(self, rows):
        self.session = types.SimpleNamespace(close=lambda: None)
        self.rows = list(rows)
        self.calls: list = []

    def get_store(self, store_id):
        return list(self.rows)

    def update_row(self, store_id, values):
        self.calls.append(("update_row", store_id, values))
        self.rows = [
            values if str(r.get("No")) == str(values.get("No")) else r
            for r in self.rows
        ]
        return {"Done": True}


def _seed_proposal(c, *, job="a", match="102", id_value=2, old="Bob", new="BOB"):
    rid = c.state.start_run(job=job, dry_run=False)
    return c.state.record_proposal(
        rid,
        job=job,
        match_value=match,
        id_value=id_value,
        payload={"No": id_value, "SbjNum": match, "I7_1_2": new},
        diff=[{"col": "I7_1_2", "old": old, "new": new}],
    )


def test_load_proposals_fills_tab(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    _seed_proposal(c)
    c._open_job("a")
    assert w.job_detail.proposals_model.rowCount() == 1
    c.stop()
    w.close()


def test_apply_proposals_writes_and_marks_applied(tmp_path, monkeypatch):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    pid = _seed_proposal(c)
    c._detail_name = "a"
    c._open_job("a")

    client = _ProposalClient([{"No": 2, "SbjNum": "102", "I7_1_2": "Bob"}])
    monkeypatch.setattr(c, "_client", lambda: client)

    c._apply_proposals([pid])
    _drain(c, "_apply_thread")

    assert len(client.calls) == 1
    assert client.calls[0][2]["I7_1_2"] == "BOB"
    assert c.state.get_proposal(pid)["status"] == "applied"
    assert w.job_detail.proposals_model.rowCount() == 0  # no longer pending
    c.stop()
    w.close()


def test_apply_stale_proposal_skips_write(tmp_path, monkeypatch):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    pid = _seed_proposal(c, old="Bob")
    c._detail_name = "a"
    c._open_job("a")

    # store cell manually changed → the recorded `old` no longer matches
    client = _ProposalClient([{"No": 2, "SbjNum": "102", "I7_1_2": "manual"}])
    monkeypatch.setattr(c, "_client", lambda: client)

    c._apply_proposals([pid])
    _drain(c, "_apply_thread")

    assert client.calls == []  # never wrote
    assert c.state.get_proposal(pid)["status"] == "stale"
    c.stop()
    w.close()


def test_dismiss_proposals_marks_dismissed_without_write(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    pid = _seed_proposal(c)
    c._detail_name = "a"
    c._open_job("a")

    c._dismiss_proposals([pid])
    assert c.state.get_proposal(pid)["status"] == "dismissed"
    assert w.job_detail.proposals_model.rowCount() == 0
    c.stop()
    w.close()


def test_job_from_detail_has_no_delete_policy(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._detail_name = "a"
    w.job_detail.load(_entry(), a_job())
    job = c._job_from_detail()
    assert not hasattr(job.write, "delete_policy")
    c.stop()
    w.close()


# -- pairing: fetch the survey XML once, fill pull terms (spec) -------------

def test_pair_survey_fills_terms_off_thread(tmp_path, monkeypatch):
    import store_sync.gui as gui
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._detail_name = "a"
    w.job_detail.load(_entry(), a_job())

    class Client:
        session = types.SimpleNamespace(close=lambda: None)
        def get_survey_xml(self, survey_id):
            return ("<Survey><SQ><QuestionID>762</QuestionID>"
                    "<VariableName>QPM1</VariableName></SQ>"
                    "<STo><QuestionID>762</QuestionID>"
                    "<VariableName>T_QPM1_1</VariableName></STo></Survey>")

    monkeypatch.setattr(c, "_client", lambda: Client())
    c._pair_survey()
    _drain(c, "_pair_thread")
    assert w.job_detail.binding_model.columns()[0].pull_term == "QPM1"
    assert w.job_detail.config_fields()["paired_survey_id"] == "SURV"
    c.stop(); w.close()


def test_pair_survey_reports_unresolved(tmp_path, monkeypatch):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._detail_name = "a"
    w.job_detail.load(_entry(), a_job())  # column var T_QPM1_1

    class Client:
        session = types.SimpleNamespace(close=lambda: None)
        def get_survey_xml(self, survey_id):
            return "<Survey/>"  # nothing resolves

    monkeypatch.setattr(c, "_client", lambda: Client())
    c._pair_survey()
    _drain(c, "_pair_thread")
    assert "unresolved" in w.status_label.text().lower()
    assert w.job_detail.binding_model.columns()[0].pull_term == ""
    c.stop(); w.close()


def test_job_from_detail_carries_pairing(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._detail_name = "a"
    w.job_detail.load(_entry(), a_job())
    w.job_detail.set_pairing("SURV", {"T_QPM1_1": "QPM1"}, {"T_QPM1_1": "QPM1"})
    job = c._job_from_detail()
    assert job.source.paired_survey_id == "SURV"
    assert job.source.survey_index == {"T_QPM1_1": "QPM1"}
    assert job.binding.columns[0].pull_term == "QPM1"
    c.stop(); w.close()


# -- blacklist: manual add / remove (spec) ----------------------------------

def test_blacklist_manual_add_and_remove_persist(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    c._detail_name = "a"
    w.blacklist_add_requested.emit("104")
    assert c.state.blacklisted_keys("a") == {"104"}
    w.blacklist_remove_requested.emit(["104"])
    assert c.state.blacklisted_keys("a") == set()
    c.stop()
    w.close()


def test_open_job_fills_blacklist_tab(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule(""), enabled=False)
    c.state.blacklist_key("a", "104", reason="deleted")
    c._open_job("a")
    assert w.job_detail.blacklist_model.rowCount() == 1
    c.stop()
    w.close()


def test_on_due_claims_minute_marker(tmp_path):
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c._on_due("a")
    # a second call in the same minute is a no-op (marker already claimed)
    assert c.state.claim_due("a", c._minute_key()) is False
    c.stop(); w.close()


def test_auto_sync_swallows_os_error(tmp_path, monkeypatch):
    from store_sync import gui
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    c.save("a", a_job(), schedule=Schedule("minute", 10), enabled=True)
    monkeypatch.setattr(
        gui.os_scheduler, "add",
        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("no schtasks")),
    )
    # must not raise
    c._auto_sync_os("a")
    c.stop(); w.close()


def test_sync_os_schedule_reports_status(tmp_path, monkeypatch):
    from store_sync import gui
    w = MainWindow()
    c = Controller(w, db_path=tmp_path / "s.db", run_fn=lambda name, stop: None)
    monkeypatch.setattr(
        gui.os_scheduler, "sync", lambda entries: "sync: 0 added, 0 removed"
    )
    c.sync_os_schedule()   # must not raise
    c.stop(); w.close()
