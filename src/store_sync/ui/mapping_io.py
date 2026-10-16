"""Excel import/export of the column mapping (spec §7.3).

The operator downloads a template (store columns pre-filled) and fills the
source-variable column. Import validates everything before it reaches a job:
unknown store columns / source vars, duplicates, and a bad header row are all
rejected with a report. A valid import is persisted through the normal Save path.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from store_sync.core.config import ColumnBinding

HEADERS = ["store_col", "source_var"]


def export_template(path: str | Path, store_cols: list[str]) -> None:
    """Write a template with the store columns pre-filled and blanks to fill.

    No `kind` column: the STG API exposes no store column type (Form1.cs
    REST surface has no `GetStoreColumns`), so a kind column was always a
    manual guess — removed rather than confuse the operator.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "mapping"
    ws.append(HEADERS)
    for col in store_cols:
        ws.append([col, ""])
    wb.save(path)


def write_mapping(path: str | Path, columns: list[ColumnBinding]) -> None:
    """Write a filled mapping (used for tests / export of an existing job)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "mapping"
    ws.append(HEADERS)
    for c in columns:
        ws.append([c.store_col, c.source_var])
    wb.save(path)


def parse_mapping(
    path: str | Path,
    valid_store_cols: list[str],
    valid_source_vars: list[str] | None = None,
    *,
    id_column: str | None = None,
) -> tuple[list[ColumnBinding], list[str]]:
    """Parse + validate a mapping workbook.

    Returns `(columns, errors)`. `errors` non-empty means the import must be
    rejected; `columns` holds only the rows that parsed cleanly.

    `valid_source_vars=None` skips source-var validation: the var list comes
    from the survey XML (a 30 s+ download), so the fast import path defers it
    to the run's readiness gate, which validates `source_var` presence against a
    live pull before any write.
    An empty list (not None) still validates and rejects everything.

    `id_column` names the store's id/tag column. Its row legitimately has no
    `source_var` — the planner owns that column (adds assign `max(id)+1`,
    updates keep the value immutable), so the row is skipped instead of being
    rejected as incomplete. Only a *known* id column row is skipped this way.
    """
    errors: list[str] = []
    wb = load_workbook(path, data_only=True)
    ws = wb.active

    header = [c.value for c in ws[1]]
    if [h for h in header[:2]] != HEADERS[:2]:
        return [], [f"bad header {header!r}; expected {HEADERS!r}"]

    store_set = set(valid_store_cols)
    source_set = set(valid_source_vars) if valid_source_vars is not None else None

    columns: list[ColumnBinding] = []
    seen_store: set[str] = set()
    seen_source: set[str] = set()

    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        store_col = _cell(row[0])
        source_var = _cell(row[1])
        if not store_col and not source_var:
            continue  # blank template row
        if not store_col or not source_var:
            if id_column and store_col == id_column and store_col in store_set:
                continue  # id/tag column: planner-owned, never mapped
            errors.append(f"incomplete row {n}: store_col/source_var both required")
            continue

        if store_col not in store_set:
            errors.append(f"unknown store column {store_col!r}")
            continue
        if source_set is not None and source_var not in source_set:
            errors.append(f"unknown source variable {source_var!r}")
            continue
        if store_col in seen_store:
            errors.append(f"duplicate store column {store_col!r}")
            continue
        if source_var in seen_source:
            errors.append(f"duplicate source variable {source_var!r}")
            continue

        seen_store.add(store_col)
        seen_source.add(source_var)
        columns.append(ColumnBinding(source_var, store_col))

    return columns, errors


def _cell(value) -> str:
    return "" if value is None else str(value).strip()
