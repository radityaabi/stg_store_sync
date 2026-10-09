"""Tests for store_sync.ui.mapping_io — Excel template + import/validate."""
from __future__ import annotations

import pytest
from openpyxl import Workbook, load_workbook

from store_sync.core.config import ColumnBinding
from store_sync.ui.mapping_io import (
    HEADERS,
    export_template,
    parse_mapping,
    write_mapping,
)


def store_cols():
    return ["I7_1_2", "I8a_1", "Q32", "Spare"]


def source_vars():
    return ["T_QPM1_1", "T_QPM1_3", "Q32", "Q42_TEXT"]


class TestTemplate:
    def test_export_template_has_headers_and_prefilled_store_cols(self, tmp_path):
        p = tmp_path / "map.xlsx"
        export_template(p, store_cols())
        wb = load_workbook(p)
        ws = wb.active
        header = [c.value for c in ws[1]]
        assert header == HEADERS
        col_a = [ws.cell(row=r, column=1).value for r in range(2, 2 + len(store_cols()))]
        assert col_a == store_cols()

    def test_round_trip_write_then_parse(self, tmp_path):
        p = tmp_path / "map.xlsx"
        cols = [
            ColumnBinding("T_QPM1_1", "I7_1_2"),
            ColumnBinding("Q32", "Q32"),
        ]
        write_mapping(p, cols)
        parsed, errors = parse_mapping(p, store_cols(), source_vars())
        assert errors == []
        assert parsed == cols


class TestParse:
    def _write(self, tmp_path, rows):
        p = tmp_path / "m.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(HEADERS)
        for r in rows:
            ws.append(r)
        wb.save(p)
        return p

    def test_valid_rows_parsed(self, tmp_path):
        p = self._write(tmp_path, [["I7_1_2", "T_QPM1_1"], ["Q32", "Q32"]])
        parsed, errors = parse_mapping(p, store_cols(), source_vars())
        assert errors == []
        assert parsed[0].store_col == "I7_1_2"
        assert parsed[0].source_var == "T_QPM1_1"

    def test_blank_rows_skipped(self, tmp_path):
        p = self._write(tmp_path, [["I7_1_2", "T_QPM1_1"], [None, None], ["Q32", "Q32"]])
        parsed, errors = parse_mapping(p, store_cols(), source_vars())
        assert len(parsed) == 2
        assert errors == []

    def test_half_filled_row_is_rejected(self, tmp_path):
        from openpyxl import Workbook
        from store_sync.ui.mapping_io import HEADERS, parse_mapping

        p = tmp_path / "m.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(HEADERS)
        ws.append(["A", "", "text"])  # store col only -> incomplete
        wb.save(p)
        cols, errors = parse_mapping(p, ["A"], ["V"])
        assert any("incomplete row" in e for e in errors)

    def test_blank_source_var_rejected_as_incomplete(self, tmp_path):
        p = self._write(tmp_path, [["I7_1_2", "T_QPM1_1"], ["I8a_1", ""]])
        parsed, errors = parse_mapping(p, store_cols(), source_vars())
        assert len(parsed) == 1
        assert any("incomplete row" in e for e in errors)

    def test_unknown_store_col_rejected(self, tmp_path):
        p = self._write(tmp_path, [["NOPE", "Q32"]])
        parsed, errors = parse_mapping(p, store_cols(), source_vars())
        assert parsed == []
        assert any("NOPE" in e for e in errors)

    def test_unknown_source_var_rejected(self, tmp_path):
        p = self._write(tmp_path, [["Q32", "NOPE"]])
        parsed, errors = parse_mapping(p, store_cols(), source_vars())
        assert parsed == []
        assert any("NOPE" in e for e in errors)

    def test_duplicate_store_col_rejected(self, tmp_path):
        p = self._write(tmp_path, [["Q32", "T_QPM1_1"], ["Q32", "Q32"]])
        parsed, errors = parse_mapping(p, store_cols(), source_vars())
        assert any("duplicate" in e.lower() for e in errors)
        assert len(parsed) == 1

    def test_duplicate_source_var_rejected(self, tmp_path):
        p = self._write(tmp_path, [["Q32", "Q32"], ["I7_1_2", "Q32"]])
        parsed, errors = parse_mapping(p, store_cols(), source_vars())
        assert any("duplicate" in e.lower() for e in errors)

    def test_legacy_kind_column_ignored(self, tmp_path):
        """A legacy 3-column file (store_col · source_var · kind) still imports.

        The kind column was removed (no API source of truth); an old template
        with it must not be rejected — the extra column is simply ignored.
        """
        p = self._write(tmp_path, [["Q32", "Q32", "number"]])
        parsed, errors = parse_mapping(p, store_cols(), source_vars())
        assert errors == []
        assert parsed[0].source_var == "Q32"
        assert parsed[0].store_col == "Q32"
        assert parsed[0].kind == "text"  # kind no longer read from the file

    def test_id_column_row_without_source_var_is_skipped(self, tmp_path):
        """The id/tag column row legitimately has no source var — skip it.

        The planner owns the id column (adds assign max(id)+1, updates keep
        the value immutable), so a template row for it is skipped, not
        rejected as incomplete.
        """
        p = self._write(tmp_path, [["No", None], ["Q32", "Q32"]])
        parsed, errors = parse_mapping(
            p, ["No", "Q32"], valid_source_vars=None, id_column="No"
        )
        assert errors == []
        assert [c.store_col for c in parsed] == ["Q32"]

    def test_other_incomplete_row_still_rejected_even_with_id_column(
        self, tmp_path
    ):
        """Only the id column's row is skipped; other half-filled rows still error."""
        p = self._write(tmp_path, [["No", None], ["Spare", ""]])
        parsed, errors = parse_mapping(
            p, ["No", "Spare"], valid_source_vars=None, id_column="No"
        )
        assert any("incomplete row" in e for e in errors)
        assert parsed == []

    def test_unknown_source_var_rejected_when_list_given(self, tmp_path):
        """With a var list supplied (non-None), unknown vars are still rejected."""
        p = self._write(tmp_path, [["Q32", "NOPE"]])
        parsed, errors = parse_mapping(p, store_cols(), ["Q32", "OTHER"])
        assert any("NOPE" in e for e in errors)
        assert parsed == []

    def test_unknown_source_var_accepted_when_list_deferred(self, tmp_path):
        """Without a var list (None), validation is deferred to preflight."""
        p = self._write(tmp_path, [["Q32", "NOPE"]])
        parsed, errors = parse_mapping(p, store_cols(), None)
        assert errors == []
        assert parsed[0].source_var == "NOPE"

    def test_missing_headers_errors(self, tmp_path):
        p = tmp_path / "bad.xlsx"
        wb = Workbook()
        wb.active.append(["foo", "bar"])
        wb.save(p)
        parsed, errors = parse_mapping(p, store_cols(), source_vars())
        assert parsed == []
        assert errors
